# -*- coding: utf-8 -*-
"""Gera o trabalho em DOCX com formatação ABNT (NBR 14724 / 6023 / 6024 / 6027 / 6028 / 10520).

Duas passadas: a 1ª gera o arquivo com sumário sem páginas, converte para PDF e
descobre em que página cada seção caiu; a 2ª regrava o sumário com as páginas reais.
"""
import copy
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import conteudo as C

FONTE = "Arial"
LARGURA_TEXTO_CM = 16.0  # 21 - 3 (esq) - 2 (dir)


# --------------------------------------------------------------------------- utils
def _run(p, texto, tam=12, negrito=False, italico=False):
    r = p.add_run(texto)
    r.font.name = FONTE
    r.font.size = Pt(tam)
    r.bold = negrito
    r.italic = italico
    # garante a fonte também para caracteres non-ASCII / east-asian
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONTE)
    return r


def par(doc, texto="", *, tam=12, negrito=False, italico=False,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo_primeira=None, recuo_esq=None,
        espaco=1.5, antes=0, depois=0, nova_pagina=False, tabs=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = espaco
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    if recuo_primeira is not None:
        pf.first_line_indent = Cm(recuo_primeira)
    if recuo_esq is not None:
        pf.left_indent = Cm(recuo_esq)
    if nova_pagina:
        pf.page_break_before = True
    if tabs:
        for pos, al, leader in tabs:
            pf.tab_stops.add_tab_stop(Cm(pos), al, leader)
    if texto:
        _run(p, texto, tam=tam, negrito=negrito, italico=italico)
    return p


def campo_pagina(paragraph):
    """Insere um campo PAGE (numeração automática) no parágrafo."""
    def _el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e

    r1 = paragraph.add_run()._element
    r1.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))

    r2 = paragraph.add_run()._element
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    r2.append(it)

    r3 = paragraph.add_run()._element
    r3.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))

    r4 = paragraph.add_run("1")
    r4.font.name = FONTE
    r4.font.size = Pt(10)

    r5 = paragraph.add_run()._element
    r5.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def config_secao(sec, *, header_dist_cm=2.0):
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(header_dist_cm)
    sec.footer_distance = Cm(2.0)


# ordem dos filhos de <w:sectPr> exigida pelo schema OOXML (CT_SectPr)
_SECTPR_ORDEM = [
    "w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr",
    "w:type", "w:pgSz", "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType",
    "w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
    "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
    "w:printerSettings", "w:sectPrChange",
]


def inicio_numeracao(sec, valor):
    """<w:pgNumType w:start="N"/> respeitando a ordem exigida do sectPr."""
    sectPr = sec._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        idx = _SECTPR_ORDEM.index("w:pgNumType")
        posteriores = {qn(t) for t in _SECTPR_ORDEM[idx + 1:]}
        destino = None
        for filho in sectPr:
            if filho.tag in posteriores:
                destino = filho
                break
        if destino is None:
            sectPr.append(pg)
        else:
            destino.addprevious(pg)
    pg.set(qn("w:start"), str(valor))


# --------------------------------------------------------------------------- doc
def montar(sumario_paginas=None):
    """sumario_paginas: dict {titulo: numero_pagina} ou None (1ª passada)."""
    doc = Document()

    est = doc.styles["Normal"]
    est.font.name = FONTE
    est.font.size = Pt(12)
    est.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)

    config_secao(doc.sections[0])

    # ---------------------------------------------------------------- CAPA
    par(doc, C.INSTITUICAO, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.5)
    for _ in range(7):
        par(doc, "", espaco=1.5)
    par(doc, C.AUTOR, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.5)
    for _ in range(6):
        par(doc, "", espaco=1.5)
    par(doc, C.TITULO, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, tam=12, espaco=1.5)
    par(doc, C.SUBTITULO, align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, espaco=1.5)
    for _ in range(8):
        par(doc, "", espaco=1.5)
    par(doc, C.CIDADE, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.0)
    par(doc, C.ANO, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.0)

    # ------------------------------------------------------- FOLHA DE ROSTO
    par(doc, C.AUTOR, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True,
        espaco=1.5, nova_pagina=True)
    for _ in range(7):
        par(doc, "", espaco=1.5)
    par(doc, C.TITULO, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, tam=12, espaco=1.5)
    par(doc, C.SUBTITULO, align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, espaco=1.5)
    for _ in range(4):
        par(doc, "", espaco=1.5)
    # nota descritiva: recuada a partir do meio da folha, espaço simples, fonte 10-12
    par(doc, C.NATUREZA, align=WD_ALIGN_PARAGRAPH.JUSTIFY, tam=12,
        recuo_esq=8.0, espaco=1.0)
    for _ in range(7):
        par(doc, "", espaco=1.5)
    par(doc, C.CIDADE, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.0)
    par(doc, C.ANO, align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True, espaco=1.0)

    # -------------------------------------------------------------- RESUMO
    par(doc, "RESUMO", align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True,
        espaco=1.5, nova_pagina=True, depois=12)
    # NBR 6028: parágrafo único, sem recuo de primeira linha, espaçamento simples
    par(doc, C.RESUMO, espaco=1.0, recuo_primeira=0, depois=12)
    p = par(doc, "", espaco=1.0)
    _run(p, "Palavras-chave: ", negrito=True)
    _run(p, C.PALAVRAS_CHAVE)

    # ------------------------------------------------------------- SUMÁRIO
    par(doc, "SUMÁRIO", align=WD_ALIGN_PARAGRAPH.CENTER, negrito=True,
        espaco=1.5, nova_pagina=True, depois=12)

    itens = [(t, txt) for (t, txt) in C.CORPO if t in ("h1", "h2")]
    itens.append(("h1", "REFERÊNCIAS"))
    for tipo, titulo in itens:
        pg = "" if sumario_paginas is None else str(sumario_paginas.get(titulo, ""))
        p = par(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, espaco=1.5, recuo_primeira=0,
                recuo_esq=(0 if tipo == "h1" else 0.6),
                tabs=[(LARGURA_TEXTO_CM - (0 if tipo == "h1" else 0.6),
                       WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)])
        _run(p, titulo, negrito=(tipo == "h1"))
        _run(p, "\t" + pg, negrito=(tipo == "h1"))

    # ------------------------------------------- SEÇÃO 2: texto com paginação
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    config_secao(sec2)
    sec2.header.is_linked_to_previous = False
    hp = sec2.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    campo_pagina(hp)
    # capa não é contada; folha de rosto = 1, resumo = 2, sumário = 3 -> texto = 4
    inicio_numeracao(sec2, 4)

    primeiro = True
    for tipo, texto in C.CORPO:
        if tipo == "h1":
            par(doc, texto, align=WD_ALIGN_PARAGRAPH.LEFT, negrito=True, espaco=1.5,
                recuo_primeira=0, depois=12, nova_pagina=not primeiro)
            primeiro = False
        elif tipo == "h2":
            par(doc, texto, align=WD_ALIGN_PARAGRAPH.LEFT, negrito=True, espaco=1.5,
                recuo_primeira=0, antes=12, depois=6)
        elif tipo == "cit":
            # NBR 10520: recuo 4 cm, fonte 10, espaçamento simples, sem aspas
            par(doc, texto, tam=10, espaco=1.0, recuo_esq=4.0, recuo_primeira=0,
                antes=6, depois=6)
        elif tipo == "p0":
            par(doc, texto, espaco=1.5, recuo_primeira=0)
        else:
            par(doc, texto, espaco=1.5, recuo_primeira=1.25)

    # --------------------------------------------------------- REFERÊNCIAS
    par(doc, "REFERÊNCIAS", align=WD_ALIGN_PARAGRAPH.LEFT, negrito=True, espaco=1.5,
        recuo_primeira=0, depois=12, nova_pagina=True)
    # NBR 6023: alinhadas à esquerda, espaço simples, separadas por linha em branco
    for ref in C.REFERENCIAS:
        par(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT, espaco=1.0, recuo_primeira=0,
            depois=12)

    return doc


# --------------------------------------------------------------------------- pdf
def para_pdf(docx_path, outdir):
    import os
    env = dict(os.environ, HOME="/tmp/lo-home")
    os.makedirs("/tmp/lo-home", exist_ok=True)
    r = subprocess.run(
        ["soffice", "--headless", "--norestore",
         "-env:UserInstallation=file:///tmp/lo-profile",
         "--convert-to", "pdf", "--outdir", outdir, docx_path],
        capture_output=True, timeout=300, env=env, text=True,
    )
    pdf = docx_path.rsplit(".", 1)[0] + ".pdf"
    if not os.path.exists(pdf):
        raise RuntimeError(
            f"conversão falhou (rc={r.returncode})\nSTDOUT: {r.stdout}\nSTDERR: {r.stderr}")


def mapear_paginas(pdf_path, itens):
    """Descobre a página ABNT de cada título. ABNT = pdf_index - 1 (capa não conta)."""
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    paginas = []
    for pg in reader.pages:
        t = (pg.extract_text() or "")
        paginas.append(re.sub(r"\s+", " ", t))

    # o sumário lista TODOS os títulos: a busca precisa começar depois dele
    inicio = 0
    for i, txt in enumerate(paginas):
        if "SUMÁRIO" in txt:
            inicio = i + 1
    if inicio == 0:
        raise RuntimeError("página do sumário não localizada no PDF")

    # os títulos aparecem em ordem no documento -> cursor monotônico
    mapa = {}
    cursor = inicio
    for titulo in itens:
        alvo = re.sub(r"\s+", " ", titulo).strip()
        for i in range(cursor, len(paginas)):
            if alvo in paginas[i]:
                mapa[titulo] = i  # índice 0-based == página ABNT (capa não conta)
                cursor = i
                break
    return mapa, len(paginas)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    docx_path = f"{out}/trabalho.docx"

    print("Passada 1: gerando documento base...")
    montar(None).save(docx_path)
    para_pdf(docx_path, out)

    titulos = [txt for (t, txt) in C.CORPO if t in ("h1", "h2")] + ["REFERÊNCIAS"]
    mapa, total = mapear_paginas(f"{out}/trabalho.pdf", titulos)
    faltando = [t for t in titulos if t not in mapa]
    print(f"  PDF com {total} páginas. Títulos localizados: {len(mapa)}/{len(titulos)}")
    if faltando:
        print("  NÃO LOCALIZADOS:", faltando)

    print("Passada 2: regravando sumário com páginas reais...")
    montar(mapa).save(docx_path)
    para_pdf(docx_path, out)

    from pypdf import PdfReader
    print(f"  Final: {len(PdfReader(f'{out}/trabalho.pdf').pages)} páginas")
    for t in titulos:
        if t in mapa:
            print(f"    {mapa[t]:>3}  {t}")
